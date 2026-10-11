from __future__ import annotations

from html import escape
from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from models import Resume


def render_html(resume: Resume) -> str:
    bits = ['<article class="resume">']
    if resume.name: bits.append(f'<h1>{escape(resume.name)}</h1>')
    if resume.contact: bits.append(f'<p class="contact">{escape(" | ".join(resume.contact.values()))}</p>')
    for section in resume.sections:
        bits.append(f'<section><h2>{escape(section.name)}</h2>')
        if section.kind == "text": bits.append(f'<p>{escape(section.content or "")}</p>')
        if section.kind == "bullets": bits.append('<ul>' + ''.join(f'<li>{escape(x.text)}</li>' for x in section.bullets or []) + '</ul>')
        if section.kind == "entries":
            for entry in section.entries or []:
                org = f' <span>{escape(entry.org)}</span>' if entry.org else ''
                dates = f'<em>{escape(entry.dates)}</em>' if entry.dates else ''
                bits.append(f'<div class="entry"><div><strong>{escape(entry.title)}</strong>{org}</div>{dates}')
                if entry.bullets: bits.append('<ul>' + ''.join(f'<li>{escape(x.text)}</li>' for x in entry.bullets) + '</ul>')
                bits.append('</div>')
        bits.append('</section>')
    return ''.join(bits) + '</article>'


def render_docx(resume: Resume) -> bytes:
    doc = Document(); sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.65); sec.left_margin = sec.right_margin = Inches(.75)
    normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(10)
    if resume.name:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(resume.name); r.bold = True; r.font.size = Pt(18)
    if resume.contact:
        p = doc.add_paragraph(' | '.join(resume.contact.values())); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
    for section in resume.sections:
        h = doc.add_paragraph(); h.paragraph_format.space_before = Pt(7); h.paragraph_format.space_after = Pt(2)
        r = h.add_run(section.name.upper()); r.bold = True; r.font.size = Pt(11)
        if section.kind == 'text': doc.add_paragraph(section.content or '')
        if section.kind == 'bullets':
            for bullet in section.bullets or []: doc.add_paragraph(bullet.text, style='List Bullet')
        if section.kind == 'entries':
            for entry in section.entries or []:
                p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(0)
                p.add_run(entry.title).bold = True
                if entry.org: p.add_run(f' | {entry.org}')
                if entry.dates: p.add_run(f'    {entry.dates}').italic = True
                for bullet in entry.bullets: doc.add_paragraph(bullet.text, style='List Bullet')
    stream = BytesIO(); doc.save(stream); return stream.getvalue()
